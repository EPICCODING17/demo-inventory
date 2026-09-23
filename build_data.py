from pathlib import Path
import openpyxl,json,re,datetime,collections
root=Path(__file__).parent
actual=next((root/'../upload').glob('Demo Unit Actual*.xlsx'))
recheck=next((root/'../upload').glob('Demo List recheck*.xlsx'))
def clean(v):
 if v is None:return ''
 if isinstance(v,(datetime.datetime,datetime.date)):return v.strftime('%Y-%m-%d')
 return re.sub(r'\s+',' ',str(v).replace('\x99','')).strip()
def status(v):
 t=clean(v).upper()
 if t in ('Ο','O','○','◯','OK','GOOD'):return 'OK'
 if t=='NG' or t.startswith('NG '):return 'NG'
 return ''
def section(name):
 m=re.search(r'(20\d\d)',name)
 if 'Recall' in name:return 'LA-7500 Recall'
 if 'Hand' in name:return 'Hand Carry'
 return m.group(1) if m else name
w=openpyxl.load_workbook(actual,read_only=True,data_only=True)
groups=[];index={}
for sh in w:
 yr=section(sh.title)
 for rownum,row in enumerate(sh.iter_rows(min_row=4,values_only=True),4):
  a=list(row)+[None]*14
  no=a[1]; demo=clean(a[2]); name=clean(a[3]); model=clean(a[4]); serial=clean(a[5]); asset=clean(a[0]); st=status(a[12]); remark=clean(a[13]); qty=clean(a[6])
  if isinstance(no,int) and (name or demo) and not (yr in ('2024','2025') and not name and not demo):
   g={'id':f'{yr}-{no}-{rownum}','year':yr,'number':str(no),'name':name or model or demo,'demo':demo if demo!='-' else '', 'asset':asset if asset!='-' else '', 'model':model,'serial':serial if serial!='-' else '', 'qty':qty,'status':st,'note':clean(a[7]),'remark':remark,'owner':'','location':'','items':[],'checkCount':0,'checkNg':0,'source':f'{sh.title}!{rownum}'}
   groups.append(g);index[(yr,no)]=g
  elif isinstance(no,(float,str)) and re.match(r'^\d+\.\d+$',clean(no)):
   n=int(float(no));g=index.get((yr,n))
   if g and (name or model):g['items'].append({'no':clean(no),'name':name or model,'model':model,'serial':serial if serial!='-' else '', 'asset':asset if asset!='-' else '', 'qty':qty,'status':st,'remark':remark,'owner':'','location':'','source':f'{sh.title}!{rownum}'})
check=openpyxl.load_workbook(recheck,read_only=True,data_only=True)['Check date OK_NG']
yr='';current=None;last_no=None;checks=0;orphans=[]
for rn,row in enumerate(check.iter_rows(values_only=True),1):
 a=list(row)+[None]*13
 marker=clean(a[0]); title=clean(a[2]); first=clean(a[1])
 if marker.startswith('Demo Unit Actual Sheet'):
  yr=section(marker);current=None;last_no=None;continue
 no=a[0]
 if isinstance(no,int) and not clean(a[4]) and not clean(a[5]):
  current=index.get((yr,no));last_no=no;continue
 isdetail=bool(clean(a[4]) or clean(a[5])) and isinstance(no,(int,float,str)) and bool(re.match(r'^\d+(?:\.\d+)?$',clean(no)))
 if not isdetail:continue
 base=int(float(no))
 g=index.get((yr,base))
 if not g:orphans.append((yr,rn,base));continue
 recent=[]
 for v in a[8:12]:
  if clean(v):recent.append(clean(v))
 check_status=next((status(v) for v in reversed(recent) if status(v)),'')
 if check_status=='NG':g['checkNg']+=1
 if check_status:g['checkCount']+=1
 if not g['owner'] and clean(a[2]):g['owner']=clean(a[2])
 if not g['location'] and clean(a[3]):g['location']=clean(a[3])
 # Numbered subitems map first; main item may have integer No. and carry own check row.
 sub=next((x for x in g['items'] if x['no']==clean(no)),None)
 if sub is None and isinstance(no,int):sub=next((x for x in g['items'] if x['serial'] and x['serial']==clean(a[6])),None)
 if sub:
  sub['checkStatus']=check_status;sub['checkRaw']=recent[-1] if recent else '';sub['owner']=clean(a[2]);sub['location']=clean(a[3]);sub['checkSource']=f'Check date OK_NG!{rn}'
  if not sub['model']:sub['model']=clean(a[5]) if clean(a[5])!='-' else ''
  if not sub['serial']:sub['serial']=clean(a[6]) if clean(a[6])!='-' else ''
  if not sub['asset']:sub['asset']=first if first!='-' else ''
 else:
  if not g['serial'] and clean(a[6])!='-':g['serial']=clean(a[6])
  if not g['model'] and clean(a[5])!='-':g['model']=clean(a[5])
  if not g['asset'] and first!='-':g['asset']=first
  g['checkStatus']=check_status;g['checkRaw']=recent[-1] if recent else ''
 checks+=1
out={'meta':{'actual':actual.name,'recheck':recheck.name,'notice':'รายการจากไฟล์ Excel ปี 2025; วันที่ตรวจในชีตอาจเก่ากว่าวันที่ไฟล์','groups':len(groups),'items':sum(len(g['items']) for g in groups)},'groups':groups}
payload=json.dumps(out,ensure_ascii=False,separators=(',',':')).replace('</','<\\/')
(root/'data/inventory.json').write_text(payload,encoding='utf-8')
(root/'data/inventory.js').write_text('window.DEMO_INVENTORY = '+payload+';\n',encoding='utf-8')
print('Groups',len(groups),'Items',out['meta']['items'],'Matched check rows',checks,'Unmatched',len(orphans),'Years',dict(collections.Counter(g['year'] for g in groups)))
print('Unmatched examples',orphans[:12])
